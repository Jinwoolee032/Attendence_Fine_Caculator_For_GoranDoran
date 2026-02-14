from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from typing import Dict, List, Any
import sys
import os

# Ensure project root is in path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from src.models import Session, ActivityType

class Exporter:
    def __init__(self):
        pass

    def export_to_excel(self, filepath: str, results: Dict[str, Dict], sessions: List[Session], deduct_fp: bool = False):
        wb = Workbook()
        ws = wb.active
        ws.title = "벌금 계산"
        
        sorted_sessions = sorted(sessions, key=lambda s: s.date)
        
        # 1. Define Columns based on Sessions and Activities
        # List of (Session, ActivityType, HeaderString)
        columns = [] 
        for s in sorted_sessions:
            if ActivityType.ATTENDANCE in s.activities:
                columns.append((s, ActivityType.ATTENDANCE, f"{s.date} 출석"))
            if ActivityType.COMMENT in s.activities:
                columns.append((s, ActivityType.COMMENT, f"{s.date} 댓글"))
                
        # Column Indices (1-based)
        # 1: Label
        # 2 ~ 1+len(columns): Data
        # Summary Starts after Data
        summary_start_col = 2 + len(columns)
        
        # 2. Write Main Headers
        ws.cell(row=1, column=1, value="이름")
        for idx, col_info in enumerate(columns):
            ws.cell(row=1, column=2+idx, value=col_info[2])
            
        # Summary Headers
        ws.cell(row=1, column=summary_start_col, value="잔여 프리패스 공제액")
        ws.cell(row=1, column=summary_start_col+1, value="총합 댓글 불참수")
        ws.cell(row=1, column=summary_start_col+2, value="댓글 벌금 부과대상")
        
        if deduct_fp:
            ws.cell(row=1, column=summary_start_col+3, value="임시벌금 합")
            ws.cell(row=1, column=summary_start_col+4, value="실납부액")
            last_col = summary_start_col+4
        else:
            ws.cell(row=1, column=summary_start_col+3, value="벌금 합")
            last_col = summary_start_col+3
        
        # Style Header
        for col in range(1, last_col+1):
            cell = ws.cell(row=1, column=col)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = Font(bold=True)
            cell.border = Border(bottom=Side(style='thick'))

        current_row = 2
        
        # Styles
        thin = Side(style='thin')
        thick = Side(style='thick')
        border_all_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
        
        for name, data in results.items():
            r = current_row
            
            # Initialize block with "-"
            init_max_col = summary_start_col+5 if deduct_fp else summary_start_col+4
            for row_offset in range(6):
                for col in range(1, init_max_col):
                    ws.cell(row=r+row_offset, column=col, value="-")

            # --- Row 1 (r): Name | Status | Debate Info ---
            ws.cell(row=r, column=1, value=name)
            
            # Fill Session Data (Status)
            for idx, (sess, act, _) in enumerate(columns):
                val = "-"
                # Determine status from data details
                if act == ActivityType.ATTENDANCE:
                    # Filter attendance_details for this session date
                    # Note: We match by date. If multiple ATT sessions per day, we might need title matching too.
                    # Current logic matches by date in calc_engine.
                    for d in data['attendance_details']:
                        if d['date'] == sess.date:
                            val = d['status']
                            break
                elif act == ActivityType.COMMENT:
                     for d in data['comment_details']:
                        if d['date'] == sess.date and d['title'] == sess.title: # Title match if possible
                            val = d['status']
                            break
                        elif d['date'] == sess.date: # Fallback
                             val = d['status']
                ws.cell(row=r, column=2+idx, value=val)

            # Debate Summary
            ws.cell(row=r, column=summary_start_col, value="토론대회 참여 여부")
            # Logic: If initial_fp has +1.0 bonus. 
            # 1.5(Reg) or 1.0(Vac). If >= 2.0, yes.
            debate_yes = data['initial_fp'] >= 2.0
            ws.cell(row=r, column=summary_start_col+1, value="예" if debate_yes else "아니오")

            # --- Row 2 (r+1): Remaining FP ---
            ws.cell(row=r+1, column=1, value="잔여 프리패스")
            for idx, (sess, act, _) in enumerate(columns):
                if act == ActivityType.ATTENDANCE:
                    val = "-"
                    for d in data['attendance_details']:
                        if d['date'] == sess.date:
                            val = d['remaining_fp']
                            break
                    ws.cell(row=r+1, column=2+idx, value=val)
            
            ws.cell(row=r+1, column=summary_start_col, value=data['remaining_fp_value'])

            # --- Row 3 (r+2): Comment Points ---
            ws.cell(row=r+2, column=1, value="댓글 불참수")
            for idx, (sess, act, _) in enumerate(columns):
                if act == ActivityType.COMMENT:
                    val = "-"
                    for d in data['comment_details']:
                        if d['date'] == sess.date and d.get('title') == sess.title:
                            val = d['point']
                            break
                        elif d['date'] == sess.date:
                            val = d['point']
                    ws.cell(row=r+2, column=2+idx, value=val)

            ws.cell(row=r+2, column=summary_start_col+1, value=data['total_comment_points'])
            ws.cell(row=r+2, column=summary_start_col+2, value="예" if data['is_comment_fine_target'] else "아니오")

            # --- Row 4 (r+3): Comment Fine ---
            ws.cell(row=r+3, column=1, value="댓글 벌금")
            for idx, (sess, act, _) in enumerate(columns):
                if act == ActivityType.COMMENT:
                    val = "-"
                    for d in data['comment_details']:
                         if d['date'] == sess.date and d.get('title') == sess.title:
                            val = d['potential_fine']
                            break
                         elif d['date'] == sess.date:
                             val = d['potential_fine']
                    ws.cell(row=r+3, column=2+idx, value=val)

            ws.cell(row=r+3, column=summary_start_col+3, value=data['total_comment_fine'])

            # --- Row 5 (r+4): Attendance Fine ---
            ws.cell(row=r+4, column=1, value="출석 벌금")
            for idx, (sess, act, _) in enumerate(columns):
                if act == ActivityType.ATTENDANCE:
                    val = "-"
                    for d in data['attendance_details']:
                        if d['date'] == sess.date:
                            val = d['fine']
                            break
                    ws.cell(row=r+4, column=2+idx, value=val)
            
            ws.cell(row=r+4, column=summary_start_col+3, value=data['total_attendance_fine'])

            # --- Row 6 (r+5): Total Fine ---
            ws.cell(row=r+5, column=1, value="벌금 총액")
            
            if deduct_fp:
                 # Logic: Actual = Max(0, AttFine - FP_Val) + CommFine
                att_fine = data['total_attendance_fine']
                comm_fine = data['total_comment_fine']
                fp_val = data['remaining_fp_value']
                deducted_att_fine = max(0, att_fine - fp_val)
                actual_payment = deducted_att_fine + comm_fine
                
                ws.cell(row=r+5, column=summary_start_col+3, value=data['total_fine'])
                ws.cell(row=r+5, column=summary_start_col+4, value=actual_payment)
            else:
                ws.cell(row=r+5, column=summary_start_col+3, value=data['total_fine'])

            # --- Styling ---
            last_col_idx = summary_start_col+4 if deduct_fp else summary_start_col+3
            
            # 1. Basic Thin Border for all cells in block
            for row in range(r, r+6):
                for col in range(1, last_col_idx+1):
                    cell = ws.cell(row=row, column=col)
                    cell.border = border_all_thin
                    cell.alignment = Alignment(horizontal='center', vertical='center')

            # 2. Thick Outer Border for Block
            self._apply_outer_border(ws, r, 1, r+5, last_col_idx, thick)

            # 3. Thick Border for Debate Info (Row 1, Cols L-M equivalent -> summary_start_col ~ +1)
            self._apply_outer_border(ws, r, summary_start_col, r, summary_start_col+1, thick)

            # 4. Thick Border for Cmt Summary (Row 3, Cols M-N equivalent -> summary_start_col+1 ~ +2)
            self._apply_outer_border(ws, r+2, summary_start_col+1, r+2, summary_start_col+2, thick)

            current_row += 6

        # Auto-fit columns
        for col_cells in ws.columns:
            length = 0
            for cell in col_cells:
                if cell.value:
                    lines = str(cell.value).split('\n')
                    longest = max(len(l) for l in lines)
                    if longest > length: length = longest
            # rough estimate
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = length * 1.5 + 2

        wb.save(filepath)

    def _apply_outer_border(self, ws, min_row, min_col, max_row, max_col, side_style):
        """Apply thick border to the outer edge of a range."""
        for row in range(min_row, max_row+1):
            for col in range(min_col, max_col+1):
                cell = ws.cell(row=row, column=col)
                # Get current border or create new
                current_border = cell.border
                # Create mutable copy args
                kwargs = {
                    'left': current_border.left,
                    'right': current_border.right,
                    'top': current_border.top,
                    'bottom': current_border.bottom
                }
                
                if row == min_row: kwargs['top'] = side_style
                if row == max_row: kwargs['bottom'] = side_style
                if col == min_col: kwargs['left'] = side_style
                if col == max_col: kwargs['right'] = side_style
                
                cell.border = Border(**kwargs)
